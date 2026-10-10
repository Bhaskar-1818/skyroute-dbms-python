"""Builds docs/SkyRoute_Slides.pptx (11 slides). Run from project root:
python docs/build_slides.py"""
from pptx import Presentation
from pptx.util import Pt

prs = Presentation()
prs.slide_width, prs.slide_height = Pt(960), Pt(540)
blank = prs.slide_layouts[6]

def slide(title, points):
    s = prs.slides.add_slide(blank)
    tx = s.shapes.add_textbox(Pt(40), Pt(20), Pt(880), Pt(80)).text_frame
    tx.text = title
    tx.paragraphs[0].runs[0].font.size = Pt(32)
    tx.paragraphs[0].runs[0].font.bold = True
    body = s.shapes.add_textbox(Pt(40), Pt(120), Pt(880), Pt(380)).text_frame
    body.word_wrap = True
    for i, pt in enumerate(points):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = "\u2022  " + pt
        p.level = 0
        for r in p.runs:
            r.font.size = Pt(20)
    return s

slide("SkyRoute: Airline Ticket & Baggage System",
      ["Multi-airport DBMS mini project — MySQL 8.0, database skyroute_db",
       "5 airports, 3 airlines, 140 dated departures, ~14,000 bookings",
       "Every manual-system failure maps to one database mechanism"])
slide("Problem statement",
      ["Duplicate seat booking by concurrent agents",
       "Gate assignment clashes between flights",
       "Baggage fee miscalculated above 20 kg",
       "No single database across stations; overselling past capacity"])
slide("ER model",
      ["Masters: Airport, Airline, Aircraft, Terminal/Gate/Runway",
       "Flight (timetable) 1—N Flight_Instance (dated departure)",
       "Passenger 1—N Booking N—1 Instance; Baggage weak entity",
       "Employee supertype with Pilot / Crew / Ground subtypes"])
slide("Schema highlights",
      ["18 InnoDB tables, surrogate integer PKs, business UNIQUE keys",
       "Flights carry departure AND arrival airports, CHECK dep <> arr",
       "UNIQUE(instance, seat): double booking impossible in storage",
       "UNIQUE(terminal, gate): gate clash impossible in storage"])
slide("Scale by design + generator",
      ["Flight/Instance split holds unlimited dated departures",
       "Deterministic generator (seed 42): 14 days, 35–80% load factors",
       "Machine-verified: zero duplicate seats, all FKs in range",
       "Knobs: NUM_DAYS, ROUTES, load range — volume on demand"])
slide("Advanced SQL",
      ["flight_board_view: whole departure board in one query",
       "booking_receipt_view: price + baggage fee (20 kg free, Rs 500/kg)",
       "trg_booking_capacity: trigger refuses overselling past capacity",
       "MakeBooking procedure: seat + payment atomically or not at all"])
slide("Transactions (ACID, demonstrated live)",
      ["Demo A: booking + payment commit together",
       "Demo B: failed payment rolls the booking back, seat stays free",
       "Demo C: two sessions show REPEATABLE READ isolation"])
slide("Indexing + cost",
      ["No re-indexing of auto-indexed FKs — only workload filters",
       "4 indexes: status, departure date, surname, bag weight",
       "Before/after EXPLAIN pairs captured in the test log"])
slide("Access control",
      ["sky_admin: all privileges; sky_staff: read-write; sky_viewer: read-only",
       "Grants verified with SHOW GRANTS on the exact database name",
       "Viewer INSERT denied — least privilege proven, not claimed"])
slide("Testing",
      ["30 checks in 8 groups (docs/TESTING.md): joins, analytics,",
       "constraint rejections, atomicity, demos, EXPLAIN pairs, grants, scale",
       "Every check states its expected result before the run"])
slide("Conclusion + next step",
      ["All 5 problem items closed by named mechanisms, all tested",
       "Flask + sqlite3 web project reuses this design with full CRUD",
       "Demo: departure board, live booking, duplicate-seat refusal, fee receipt"])

prs.save("docs/SkyRoute_Slides.pptx")
print("slides written")
