"""Builds dbms/docs/SkyRoute_DBMS_Report.docx from REPORT_TEXT.md.
Run from repo root:  python dbms/docs/build_report.py"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from docx import Document  # noqa: E402
from docx.shared import Pt  # noqa: E402

from report_content import SCHEMA_TABLE, load  # noqa: E402

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)


def table(headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for j, hd in enumerate(headers):
        t.rows[0].cells[j].text = hd
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            t.rows[i + 1].cells[j].text = str(val)
    doc.add_paragraph("")


for kind, payload in load():
    if kind == "title":
        doc.add_heading(payload, level=0)
    elif kind in ("sub", "meta"):
        doc.add_paragraph(payload)
    elif kind == "h1":
        doc.add_heading(payload, level=1)
        if payload.startswith("6."):
            headers, rows = SCHEMA_TABLE
            table(headers, rows)
    elif kind == "h2":
        doc.add_heading(payload, level=2)
    elif kind == "p":
        doc.add_paragraph(payload)
    elif kind == "bullets":
        for it in payload:
            doc.add_paragraph(it, style="List Bullet")

doc.save("dbms/docs/SkyRoute_DBMS_Report.docx")
print("report written")
