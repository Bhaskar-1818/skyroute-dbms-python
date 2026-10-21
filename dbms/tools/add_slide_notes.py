"""Adds speaker + animation notes to each slide (handoff for the animated
Claude PPT pass; python-pptx cannot author animations itself).
Run from repo root after build_slides.py."""
from pptx import Presentation

NOTES = [
    "Cover. Animation: headline words fly in one by one; plane motif draws across. Say: one database for every station.",
    "Problem. Animation: each row wipes in with its mechanism. Stress: every failure maps to exactly one mechanism.",
    "ER. Animation: entities fade in layer by layer (masters, operations, dependents). Point at the weak entity.",
    "Schema table. Animation: rows appear one by one. Pause on the seat and gate UNIQUE rows.",
    "Scale. Animation: count-up effect on 336 / 34k if available. Key line: scale is a knob, not a rewrite.",
    "Advanced SQL. Animation: each bullet with its code snippet. Demo the trigger refusal live if possible.",
    "Transactions. Animation: Demo B rollback as a two-step reveal. The seat stays free.",
    "Indexing. Animation: before/after EXPLAIN bars. Say the FK line out loud: no cargo cult.",
    "Access + testing. Animation: grant table, then the denied INSERT. 30 checks, all expected-first.",
    "Flask site. Animation: login gate, manifest table reveal. Live demo: traveler books, staff checks in from the manifest.",
    "Result. Animation: checklist ticks. Close with the demo invitation.",
]

prs = Presentation("dbms/docs/report/SkyRoute_Slides.pptx")
for i, slide in enumerate(prs.slides):
    slide.notes_slide.placeholders[1].text = NOTES[i] if i < len(NOTES) else ""
prs.save("dbms/docs/report/SkyRoute_Slides.pptx")
print("notes added")

